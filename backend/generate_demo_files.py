import io
from docx import Document
from reportlab.pdfgen import canvas

def create_demo_files():
    # 1. TXT File
    txt_content = """Customer Complaint Report
Customer: ABC Healthcare Ltd.
Product: Neurovit Capsules
Strength: 500 mg
Batch Number: NV24081
Manufacturing Date: 02/08/2026
Expiry Date: 01/08/2028
Quantity Affected: 20 bottles
Complaint Date: 10/09/2026
Complaint: Several capsules were found discolored inside sealed bottles."""
    
    with open("../demo_data/sample_discoloration_complaint.txt", "w", encoding="utf-8") as f:
        f.write(txt_content)
        
    # 2. PDF File
    c = canvas.Canvas("../demo_data/sample_batch_quality_complaint.pdf")
    c.drawString(100, 800, "Customer Complaint Report")
    y = 770
    for line in txt_content.split('\n')[1:]:
        c.drawString(100, y, line)
        y -= 20
    c.save()

    # 3. EML File
    eml_content = """From: jane.doe@abchealthcare.com
To: QA@pharmacomplaints.com
Subject: Customer Complaint: Neurovit Discoloration
Date: Thu, 10 Sep 2026 14:30:00 -0400

Customer Complaint Report
Customer: ABC Healthcare Ltd.
Product: Neurovit Capsules
Strength: 500 mg
Batch Number: NV24081
Manufacturing Date: 02/08/2026
Expiry Date: 01/08/2028
Quantity Affected: 20 bottles
Complaint Date: 10/09/2026
Complaint: Several capsules were found discolored inside sealed bottles."""

    with open("../demo_data/sample_customer_complaint.eml", "w", encoding="utf-8") as f:
        f.write(eml_content)
        
    # 4. DOCX File
    doc = Document()
    doc.add_heading('Customer Complaint Report', 0)
    for line in txt_content.split('\n')[1:]:
        doc.add_paragraph(line)
    doc.save("../demo_data/sample_complaint.docx")

if __name__ == "__main__":
    create_demo_files()
